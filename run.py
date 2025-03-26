import os
import glob
import subprocess
import sys
import platform
from docx import Document
import pandas as pd
import openpyxl
from pprint import pprint
import json

# Detect OS
IS_WINDOWS = platform.system() == "Windows"
IS_MAC = platform.system() == "Darwin"
# Convert .doc to .docx if necessary
def convert_doc_to_docx(doc_path):
    """Convert .doc to .docx if needed."""
    if doc_path.endswith(".docx"):
        return doc_path  # No conversion needed

    docx_path = doc_path + "x"  # Change .doc to .docx

    print(f"🔄 Converting {doc_path} to {docx_path}...")

    if os.name == "nt":  # Windows
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        doc = word.Documents.Open(doc_path)
        doc.SaveAs2(docx_path, FileFormat=16)
        doc.Close()
        word.Quit()
    else:  # macOS/Linux
        subprocess.run(["soffice", "--headless", "--convert-to", "docx", doc_path])

    if os.path.exists(docx_path):
        print("✅ Conversion successful!")
        return docx_path
    else:
        print("❌ Conversion failed!")
        exit(1)

# === CONSTANTS ===
OPPORTUNITY_OWNER = "Dave Hoffman"
ROLE_KEYS = {
    "Salesperson", "FAE",
    "Supplier Contact 1", "Supplier Contact 2",
    "Project Engineer", "Additional Engineer", "Purchasing Contact"
}
SCOPED_KEYS = {"Phone", "Email"}

def parse_docx_table(file_path):
    doc = Document(file_path)
    table = doc.tables[0]
    data = {}
    current_section = "Initial Info"
    data[current_section] = {}
    last_role = None

    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]

        if all(cell == "" for cell in cells):
            continue  # skip empty or whitespace-only rows

        # Remove consecutive duplicates (handles merged cells)
        unique_cells = []
        for cell in cells:
            if not unique_cells or cell != unique_cells[-1]:
                unique_cells.append(cell)
        print(f"Row cells: {unique_cells}")
        cells = unique_cells

        # Detect section titles
        if len(row.cells) == 1 or (len(cells) == 1 and len(row.cells) > 1):
            current_section = cells[0]
            data[current_section] = {}
            last_role = None  # reset scope when section changes
            continue

        if current_section not in data:
            data[current_section] = {}

        for i in range(0, len(cells)-1, 2):
            key, value = cells[i], cells[i+1]

            if key in ROLE_KEYS:
                last_role = key
                data[current_section][key] = value
            elif key in SCOPED_KEYS and last_role:
                scoped_key = f"{last_role} {key}"
                data[current_section][scoped_key] = value
            else:
                data[current_section][key] = value
                last_role = None # reset if not a scoped field
    return data

def flatten_dict(nested_dict, parent_key=''):
    items = {}
    for k, v in nested_dict.items():
        new_key = f"{parent_key}.{k}" if parent_key else k
        if isinstance(v, dict):
            items.update(flatten_dict(v, new_key))
        else:
            items[new_key] = v
    return items

def write_to_excel(parsed_data, excel_file_path, header_mapping, opportunity_owner):
    flat_data = flatten_dict(parsed_data)

    # Load workbook and sheet
    wb = openpyxl.load_workbook(excel_file_path)
    sheet = wb.active

    headers = [cell.value for cell in sheet[1]]
    print(f"Excel headers: {headers}")
    new_row = []

    for header in headers:
        # Find the parsed_data key associated with this header
        match = next((k for k, v in header_mapping.items() if v == header), None)
        if match:
            if match.startswith("static."):
                if match == "static.Opportunity Owner":
                    new_row.append(opportunity_owner)
                else:
                    new_row.append("")
            else:
                new_row.append(flat_data.get(match, ""))
        else:
            new_row.append("")  # For manual or unmatched columns

    print(f"Appending row: {new_row}")
    # Find the last row with actual values in it
    last_row = sheet.max_row
    while last_row > 1 and all(cell.value is None for cell in sheet[last_row]):
        last_row -= 1

    # Append below that row
    sheet.insert_rows(last_row + 1)
    for col_index, value in enumerate(new_row, start=1):
        sheet.cell(row=last_row + 1, column=col_index).value = value
    wb.save(excel_file_path)

# Get the directory of the script
script_dir = os.path.dirname(os.path.abspath(__file__))

# Define the folder containing the Word files
word_file_dir = os.path.join(script_dir, "word_file")

# Find all Word files (.doc or .docx) in the folder
word_files = [file for pattern in ["*.doc*", "*.DOC*"] for file in glob.glob(os.path.join(word_file_dir, pattern)) if not os.path.basename(file).startswith("~$")]

# Ensure there is at least one file
if not word_files:
    print("❌ Error: No Word files found in 'word_file/' folder!")
    exit(1)

# If multiple files exist, sort by last modified time (newest first)
if len(word_files) > 1:
    word_files.sort(key=os.path.getmtime, reverse=True)
    print("⚠️ Multiple Word files found. Using the most recently modified file.")

# Select the first file (either the only one or the newest)
word_file = word_files[0]
print(f"📄 Using Word file: {word_file}")

parsed_data = parse_docx_table(word_file)
print("################################################################")
print(f"script_dir: {script_dir}")
# excel_file_path = os.path.join(script_dir, "../notes on JWT reg tracker.xlsx")
excel_file_path = os.path.join(script_dir, "./excel_file/notes on JWT reg tracker.xlsx")
print(f"Resolved Excel path: {os.path.abspath(excel_file_path)}")
print("################################################################")
print(f"Parsed data: {json.dumps(parsed_data, indent=2)}")

# Mapping from parsed_data to Excel headers
excel_mapping = {
    "Initial Info.Date": "Date",
    "static.Opportunity Owner": "Opportunity Owner",
    "Design Customer Information.Name": "Account Name",
    "Design Customer Information.State": "Design State",
    "Purchasing Customer Information.State": "Mfg State / CM if Different",
    "Distributor Information.Salesperson": "Distributor DR",
    "Distributor Information.FAE": "Dist FAE",
    "Registration Information.Avnet Registration ID": "Design Reg Number",
    "Project Information.Project Name": "Program Name",
    "Project Information.Project Engineer": "Customer Contact",
    "Project Information.Project Engineer Phone": "Phone",
    "Project Information.Project Engineer Email": "Email",
    "Part Information.Supplier Part": "CTS Part Number",
    "Part Information.Annual Part Quantity": "Volume",
    "Part Information.Annual Production Revenue": "Value"
    # Skipped: "Split Submitted" and "CTS Product Line" — manual entry
}

write_to_excel(parsed_data, excel_file_path, excel_mapping, OPPORTUNITY_OWNER)
